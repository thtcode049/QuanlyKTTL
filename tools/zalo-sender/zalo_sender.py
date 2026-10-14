"""
Zalo Bulk Sender - Gửi tin nhắn hàng loạt trên chat.zalo.me theo danh sách SĐT (đã kết bạn).

Công nghệ: Python + Playwright (Sync API) + CustomTkinter (đa luồng).

Quy trình mỗi số:
  fill SĐT vào input#contact-search-input
  -> click kết quả đầu tiên div[id^='friend-item-']
  -> keyboard.insert_text() nội dung (đã trộn Spintax) vào div#richInput -> Enter

Chống Captcha: nghỉ ngẫu nhiên 2-5s giữa các số, nghỉ 30s sau mỗi 40 tin (chỉnh được trên giao diện),
tự động TẠM DỪNG khi lỗi liên tiếp 3 lần để bạn kiểm tra trình duyệt.

Chạy: Chay_Zalo_Sender.bat  (hoặc: .venv\\Scripts\\python zalo_sender.py)
"""

from __future__ import annotations

import csv
import json
import os
import queue
import random
import re
import threading
import time
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox

import customtkinter as ctk
from playwright.sync_api import Error as PWError
from playwright.sync_api import TimeoutError as PWTimeout
from playwright.sync_api import sync_playwright

APP_DIR = Path(__file__).resolve().parent
PROFILE_DIR = APP_DIR / "zalo_profile"  # Phiên đăng nhập Zalo Web được lưu tại đây
LOG_DIR = APP_DIR / "logs"
CONFIG_FILE = APP_DIR / "config.json"
ZALO_URL = "https://chat.zalo.me/"

SEL_SEARCH = "input#contact-search-input"
SEL_FRIEND = "div[id^='friend-item-']"
SEL_INPUT = "div#richInput"

MAX_CONSECUTIVE_ERRORS = 3

# ----------------------------------------------------------------------------
# Xử lý dữ liệu: chuẩn hóa SĐT + Spintax
# ----------------------------------------------------------------------------

# Tìm cụm giống SĐT VN trong 1 dòng: +84 / 84 / 0084 / 0 (hoặc mất số 0 do Excel) + 9 số đầu 3/5/7/8/9.
# Cho phép dấu cách, chấm, gạch giữa các số: "0912 345 678", "+84.912.345.678"...
PHONE_FIND_RE = re.compile(r"(?<!\d)(?:\+?84|0084|0)?[ .\-]?[35789](?:[ .\-]?\d){8}(?!\d)")
PHONE_NORMAL_RE = re.compile(r"^(?:84|0084|0)?([35789]\d{8})$")
SPIN_RE = re.compile(r"\{([^{}]*)\}")


def normalize_phone(raw: str) -> str | None:
    """Đưa SĐT về dạng 10 số bắt đầu bằng 0. Trả về None nếu không hợp lệ."""
    m = PHONE_NORMAL_RE.match(re.sub(r"\D", "", raw))
    return "0" + m.group(1) if m else None


def parse_phones(text: str) -> tuple[list[str], list[str], int]:
    """Tách SĐT từ văn bản tự do. Trả về (danh sách hợp lệ không trùng, dòng lỗi, số lượng trùng)."""
    phones: list[str] = []
    seen: set[str] = set()
    invalid: list[str] = []
    dup = 0
    for line in text.splitlines():
        if not line.strip():
            continue
        found = PHONE_FIND_RE.findall(line)
        if not found:
            if sum(c.isdigit() for c in line) >= 6:
                invalid.append(line.strip())
            continue
        for raw in found:
            p = normalize_phone(raw)
            if not p:
                invalid.append(raw.strip())
            elif p in seen:
                dup += 1
            else:
                seen.add(p)
                phones.append(p)
    return phones, invalid, dup


def spin(text: str) -> str:
    """Trộn Spintax, hỗ trợ lồng nhau: {Chào|Xin chào} {anh|chị} -> 1 biến thể ngẫu nhiên."""
    while True:
        new = SPIN_RE.sub(lambda m: random.choice(m.group(1).split("|")), text)
        if new == text:
            return new
        text = new


def read_phone_file(path: str) -> str:
    """Đọc file .txt/.csv/.xlsx và trả về văn bản thô (mỗi ô 1 dòng)."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is None:
                        continue
                    # Excel hay lưu SĐT dạng số (912345678.0) -> bỏ phần thập phân
                    lines.append(str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
        wb.close()
        return "\n".join(lines)
    for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
        try:
            return Path(path).read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return ""


@dataclass
class SendConfig:
    delay_min: float = 2.0
    delay_max: float = 5.0
    batch_size: int = 40
    batch_pause: float = 30.0
    search_timeout: float = 6.0


# ----------------------------------------------------------------------------
# Worker: toàn bộ Playwright chạy trong 1 thread duy nhất (yêu cầu của Sync API)
# ----------------------------------------------------------------------------


class ZaloWorker(threading.Thread):
    def __init__(self, emit):
        super().__init__(daemon=True)
        self.emit = emit  # emit(kind, payload) -> đẩy sự kiện về GUI
        self.cmds: queue.Queue = queue.Queue()
        self.stop_event = threading.Event()
        self.pause_event = threading.Event()  # set = đang tạm dừng
        self.pw = None
        self.ctx = None
        self.page = None
        self._ctx_closed = True

    # ---------- API cho GUI ----------
    def submit(self, cmd: str, **kw):
        self.cmds.put((cmd, kw))

    def log(self, text: str, level: str = "info"):
        self.emit("log", (level, text))

    # ---------- Vòng lặp chính ----------
    def run(self):
        try:
            with sync_playwright() as pw:
                self.pw = pw
                while True:
                    cmd, kw = self.cmds.get()
                    if cmd == "quit":
                        break
                    self.emit("busy", cmd)
                    try:
                        if cmd == "open":
                            self._ensure_logged_in()
                        elif cmd == "send":
                            self._send_batch(**kw)
                    except Exception as e:  # noqa: BLE001 - báo mọi lỗi lên GUI thay vì chết thread
                        self.log(f"Lỗi không mong muốn: {e}", "error")
                    finally:
                        self.emit("idle", None)
                self._close()
        except Exception as e:  # noqa: BLE001
            self.emit("log", ("error", f"Không khởi động được Playwright: {e}"))

    # ---------- Trình duyệt ----------
    def _launch(self):
        PROFILE_DIR.mkdir(parents=True, exist_ok=True)
        args = ["--disable-blink-features=AutomationControlled", "--window-size=1280,860"]
        last_err = None
        # Ưu tiên Chrome/Edge có sẵn trên máy -> không cần tải Chromium của Playwright
        for channel in ("chrome", "msedge", None):
            try:
                kw = dict(
                    user_data_dir=str(PROFILE_DIR),
                    headless=False,
                    args=args,
                    no_viewport=True,
                    ignore_default_args=["--enable-automation"],
                )
                if channel:
                    kw["channel"] = channel
                ctx = self.pw.chromium.launch_persistent_context(**kw)
                self.log(f"Đã mở trình duyệt: {channel or 'Chromium (Playwright)'}")
                return ctx
            except PWError as e:
                last_err = e
        raise RuntimeError(
            "Không mở được trình duyệt. Hãy đóng các cửa sổ tool cũ, cài Google Chrome, "
            f"hoặc chạy: .venv\\Scripts\\playwright install chromium\nChi tiết: {last_err}"
        )

    def _on_ctx_close(self):
        self._ctx_closed = True
        self.emit("login", False)

    def _alive(self) -> bool:
        if self.ctx is None or self._ctx_closed or self.page is None:
            return False
        try:
            return not self.page.is_closed()
        except PWError:
            return False

    def _ensure_page(self):
        if self._alive():
            return
        self._close()
        self.ctx = self._launch()
        self._ctx_closed = False
        self.ctx.on("close", lambda _: self._on_ctx_close())
        self.page = self.ctx.pages[0] if self.ctx.pages else self.ctx.new_page()
        self.page.goto(ZALO_URL, wait_until="domcontentloaded", timeout=60000)

    def _close(self):
        if self.ctx is not None and not self._ctx_closed:
            try:
                self.ctx.close()
            except PWError:
                pass
        self.ctx = None
        self.page = None
        self._ctx_closed = True

    def _ensure_logged_in(self, timeout: float = 300) -> bool:
        self._ensure_page()
        page = self.page
        if "chat.zalo.me" not in page.url:
            page.goto(ZALO_URL, wait_until="domcontentloaded", timeout=60000)
        self.emit("status", "Đang kiểm tra đăng nhập Zalo...")
        start = time.time()
        hinted = False
        while time.time() - start < timeout:
            if self.stop_event.is_set():
                return False
            try:
                if page.locator(SEL_SEARCH).first.is_visible():
                    self.emit("login", True)
                    self.emit("status", "Đã đăng nhập Zalo")
                    self.log("Đã đăng nhập Zalo Web (phiên được lưu cho lần sau).", "ok")
                    return True
            except PWError:
                pass
            if not hinted and time.time() - start > 5:
                hinted = True
                self.log("Chưa đăng nhập: mở app Zalo trên điện thoại -> Quét mã QR trên trình duyệt.", "warn")
                self.emit("status", "Đang chờ quét mã QR (tối đa 5 phút)...")
            page.wait_for_timeout(1000)
        self.log("Hết thời gian chờ đăng nhập.", "error")
        self.emit("status", "Chưa đăng nhập")
        return False

    # ---------- Điều khiển nghỉ / tạm dừng ----------
    def _checkpoint(self) -> bool:
        """Chặn khi đang tạm dừng. Trả về False nếu người dùng bấm Dừng."""
        while self.pause_event.is_set() and not self.stop_event.is_set():
            self.stop_event.wait(0.2)
        return not self.stop_event.is_set()

    def _sleep(self, seconds: float) -> bool:
        end = time.time() + seconds
        while True:
            if self.stop_event.is_set():
                return False
            if self.pause_event.is_set():
                end += 0.2  # thời gian tạm dừng không tính vào thời gian nghỉ
            remaining = end - time.time()
            if remaining <= 0:
                return True
            self.stop_event.wait(min(0.2, remaining))

    # ---------- Gửi ----------
    def _wait_friend(self, last_id: str | None, timeout: float):
        """Chờ kết quả friend-item đầu tiên. Bỏ qua nếu vẫn là người vừa gửi (kết quả cũ chưa cập nhật)."""
        page = self.page
        page.wait_for_timeout(500)  # chờ Zalo debounce ô tìm kiếm
        end = time.time() + timeout
        while time.time() < end:
            if self.stop_event.is_set():
                return None
            try:
                first = page.locator(SEL_FRIEND).first
                if first.count() and first.is_visible():
                    if first.get_attribute("id") != last_id:
                        return first
            except PWError:
                pass
            page.wait_for_timeout(150)
        return None

    def _input_cleared(self, box, timeout_ms: int) -> bool:
        end = time.time() + timeout_ms / 1000
        while time.time() < end:
            try:
                if not box.inner_text(timeout=1000).strip():
                    return True
            except PWError:
                pass
            self.page.wait_for_timeout(150)
        return False

    def _send_one(self, phone: str, msg: str, last_id: str | None, cfg: SendConfig):
        """Trả về (status, tên, chi tiết, friend_id)."""
        page = self.page
        try:
            search = page.locator(SEL_SEARCH)
            search.click(timeout=5000)
            search.fill("")
            search.fill(phone)
            item = self._wait_friend(last_id, cfg.search_timeout)
            if item is None and not self.stop_event.is_set():
                # Dự phòng: một số bản Zalo không nhận sự kiện từ fill() -> gõ từng phím
                search.fill("")
                search.press_sequentially(phone, delay=25)
                item = self._wait_friend(last_id, cfg.search_timeout / 2)
            if item is None:
                search.fill("")
                page.keyboard.press("Escape")
                return "NOT_FOUND", "", "Không thấy trong danh sách bạn bè", None

            fid = item.get_attribute("id")
            try:
                name = item.inner_text(timeout=1000).strip().split("\n")[0]
            except PWError:
                name = ""
            item.click()

            box = page.locator(SEL_INPUT)
            box.wait_for(state="visible", timeout=8000)
            box.click()
            page.keyboard.press("Control+A")
            page.keyboard.press("Delete")  # xóa bản nháp cũ nếu có
            for k, line in enumerate(msg.split("\n")):
                if k:
                    page.keyboard.press("Shift+Enter")  # xuống dòng, không gửi
                if line:
                    page.keyboard.insert_text(line)
            page.wait_for_timeout(250)
            page.keyboard.press("Enter")

            if self._input_cleared(box, 4000):
                return "OK", name, "", fid
            return "UNCONFIRMED", name, "Đã nhấn Enter nhưng ô chat chưa trống", fid
        except PWTimeout as e:
            return "ERROR", "", f"Timeout: {str(e).splitlines()[0]}", None
        except PWError as e:
            return "ERROR", "", str(e).splitlines()[0], None

    def _send_batch(self, phones: list[str], template: str, cfg: SendConfig):
        if not self._ensure_logged_in():
            return
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        log_path = LOG_DIR / f"zalo_send_{datetime.now():%Y%m%d_%H%M%S}.csv"
        total = len(phones)
        counts = {"OK": 0, "UNCONFIRMED": 0, "NOT_FOUND": 0, "ERROR": 0}
        sent_since_break = 0
        consecutive_err = 0
        last_id = None
        self.log(f"Bắt đầu gửi {total} số. Log: {log_path.name}", "info")

        with open(log_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["thoi_gian", "sdt", "ten_zalo", "trang_thai", "chi_tiet", "noi_dung"])
            for i, phone in enumerate(phones, 1):
                if not self._checkpoint():
                    break
                if not self._alive():
                    self.log("Trình duyệt đã bị đóng, đang mở lại...", "warn")
                    if not self._ensure_logged_in():
                        break
                    last_id = None

                self.emit("status", f"[{i}/{total}] Đang gửi {phone}...")
                msg = spin(template)
                status, name, detail, fid = self._send_one(phone, msg, last_id, cfg)
                if fid:
                    last_id = fid
                counts[status] += 1
                w.writerow([datetime.now().strftime("%H:%M:%S"), phone, name, status, detail, msg])
                f.flush()
                self.emit("result", {"i": i, "total": total, "phone": phone, "name": name,
                                     "status": status, "detail": detail})

                if status in ("OK", "UNCONFIRMED"):
                    sent_since_break += 1
                    consecutive_err = 0
                elif status == "ERROR":
                    consecutive_err += 1
                    if consecutive_err >= MAX_CONSECUTIVE_ERRORS:
                        consecutive_err = 0
                        self.pause_event.set()
                        self.emit("paused", True)
                        self.log(f"Lỗi {MAX_CONSECUTIVE_ERRORS} lần liên tiếp -> TỰ ĐỘNG TẠM DỪNG. "
                                 "Kiểm tra trình duyệt (Captcha / bị đăng xuất / popup) rồi bấm Tiếp tục.",
                                 "error")

                if i == total:
                    break
                if cfg.batch_size > 0 and sent_since_break >= cfg.batch_size:
                    sent_since_break = 0
                    self.emit("status", f"Đã gửi {cfg.batch_size} tin, nghỉ {cfg.batch_pause:.0f}s chống Captcha...")
                    self.log(f"Nghỉ dài {cfg.batch_pause:.0f}s sau {cfg.batch_size} tin.", "warn")
                    if not self._sleep(cfg.batch_pause):
                        break
                else:
                    d = random.uniform(cfg.delay_min, cfg.delay_max)
                    self.emit("status", f"[{i}/{total}] Nghỉ {d:.1f}s...")
                    if not self._sleep(d):
                        break

        stopped = self.stop_event.is_set()
        self.emit("done", {"counts": counts, "log": str(log_path), "stopped": stopped})


# ----------------------------------------------------------------------------
# Giao diện CustomTkinter
# ----------------------------------------------------------------------------

C_BG = "#0d1117"
C_CARD = "#161b26"
C_BORDER = "#263043"
C_ACCENT = "#0068ff"  # xanh Zalo
C_ACCENT_H = "#1a7dff"
C_OK = "#22c55e"
C_WARN = "#f59e0b"
C_ERR = "#ef4444"
C_MUTED = "#8b98ad"

STATUS_LABEL = {
    "OK": ("✅ Đã gửi", "ok"),
    "UNCONFIRMED": ("⚠️ Chưa xác nhận", "warn"),
    "NOT_FOUND": ("🔍 Không tìm thấy", "warn"),
    "ERROR": ("❌ Lỗi", "error"),
}

DEFAULT_MESSAGE = (
    "{Chào|Xin chào|Kính chào} {anh/chị|Quý phụ huynh},\n"
    "CLB Cờ vua Kiện Tướng Tương Lai {xin gửi|gửi tới} {thông báo|thông tin} mới nhất. "
    "{Cảm ơn|Trân trọng cảm ơn} {anh/chị|Quý phụ huynh}!"
)


class App(ctk.CTk):
    def __init__(self):
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")
        super().__init__(fg_color=C_BG)
        self.title("Zalo Bulk Sender • CLB Cờ vua KTTLL")
        self.geometry("1200x800")
        self.minsize(1020, 700)

        self.ui_q: queue.Queue = queue.Queue()
        self.worker = ZaloWorker(lambda k, p: self.ui_q.put((k, p)))
        self.worker.start()
        self.busy = False
        self.failed: list[str] = []
        self.counts = {k: 0 for k in STATUS_LABEL}

        self.f_title = ctk.CTkFont("Segoe UI", 22, "bold")
        self.f_h = ctk.CTkFont("Segoe UI", 14, "bold")
        self.f_body = ctk.CTkFont("Segoe UI", 13)
        self.f_mono = ctk.CTkFont("Consolas", 12)

        self._build()
        self._load_config()
        self.after(100, self._poll)
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ---------- Dựng giao diện ----------
    def _card(self, parent, **grid):
        card = ctk.CTkFrame(parent, fg_color=C_CARD, corner_radius=14, border_width=1, border_color=C_BORDER)
        card.grid(**grid)
        return card

    def _btn(self, parent, text, cmd, color=C_ACCENT, hover=C_ACCENT_H, width=120):
        return ctk.CTkButton(parent, text=text, command=cmd, fg_color=color, hover_color=hover,
                             corner_radius=10, height=36, width=width, font=self.f_h)

    def _build(self):
        self.grid_columnconfigure((0, 1), weight=1, uniform="col")
        self.grid_rowconfigure(1, weight=3)
        self.grid_rowconfigure(4, weight=2)

        # Header
        head = ctk.CTkFrame(self, fg_color="transparent")
        head.grid(row=0, column=0, columnspan=2, sticky="ew", padx=18, pady=(14, 6))
        ctk.CTkLabel(head, text="⚡ Zalo Bulk Sender", font=self.f_title, text_color="#e6edf7").pack(side="left")
        self.lbl_login = ctk.CTkLabel(head, text="●  Chưa kết nối Zalo", font=self.f_h, text_color=C_MUTED)
        self.lbl_login.pack(side="right")

        # Danh sách SĐT
        left = self._card(self, row=1, column=0, sticky="nsew", padx=(18, 8), pady=6)
        left.grid_columnconfigure(0, weight=1)
        left.grid_rowconfigure(1, weight=1)
        ctk.CTkLabel(left, text="📱 Danh sách số điện thoại (đã kết bạn)", font=self.f_h).grid(
            row=0, column=0, sticky="w", padx=14, pady=(12, 4))
        self.txt_phones = ctk.CTkTextbox(left, font=self.f_mono, fg_color=C_BG, border_width=1,
                                         border_color=C_BORDER, corner_radius=10)
        self.txt_phones.grid(row=1, column=0, sticky="nsew", padx=14)
        bar = ctk.CTkFrame(left, fg_color="transparent")
        bar.grid(row=2, column=0, sticky="ew", padx=14, pady=10)
        self._btn(bar, "📂 Mở file", self._open_file, width=105).pack(side="left")
        self._btn(bar, "✨ Chuẩn hóa", self._normalize, color="#7c3aed", hover="#8b5cf6", width=115).pack(
            side="left", padx=8)
        self._btn(bar, "🗑", lambda: self.txt_phones.delete("1.0", "end"), color="#334155",
                  hover="#475569", width=44).pack(side="left")
        self.lbl_phone_count = ctk.CTkLabel(bar, text="", font=self.f_body, text_color=C_MUTED)
        self.lbl_phone_count.pack(side="right")

        # Nội dung tin nhắn
        right = self._card(self, row=1, column=1, sticky="nsew", padx=(8, 18), pady=6)
        right.grid_columnconfigure(0, weight=1)
        right.grid_rowconfigure(1, weight=1)
        ctk.CTkLabel(right, text="💬 Nội dung tin nhắn  •  Spintax {A|B|C}", font=self.f_h).grid(
            row=0, column=0, sticky="w", padx=14, pady=(12, 4))
        self.txt_msg = ctk.CTkTextbox(right, font=self.f_body, fg_color=C_BG, border_width=1,
                                      border_color=C_BORDER, corner_radius=10, wrap="word")
        self.txt_msg.grid(row=1, column=0, sticky="nsew", padx=14)
        pv = ctk.CTkFrame(right, fg_color="transparent")
        pv.grid(row=2, column=0, sticky="ew", padx=14, pady=10)
        pv.grid_columnconfigure(1, weight=1)
        self._btn(pv, "🎲 Xem trước", self._preview, color="#0f766e", hover="#14b8a6", width=115).grid(
            row=0, column=0, sticky="nw")
        self.lbl_preview = ctk.CTkLabel(pv, text="Bấm 'Xem trước' để thử một biến thể ngẫu nhiên.",
                                        font=self.f_body, text_color=C_MUTED, justify="left",
                                        anchor="w", wraplength=420)
        self.lbl_preview.grid(row=0, column=1, sticky="ew", padx=(10, 0))

        # Cài đặt
        st = self._card(self, row=2, column=0, columnspan=2, sticky="ew", padx=18, pady=6)
        ctk.CTkLabel(st, text="⚙️ Chống Captcha", font=self.f_h).pack(side="left", padx=(14, 10), pady=10)
        self.e_cfg = {}
        for key, label, w in (("delay_min", "Nghỉ từ (s)", 50), ("delay_max", "đến (s)", 50),
                              ("batch_size", "Mỗi N tin", 50), ("batch_pause", "nghỉ (s)", 50),
                              ("search_timeout", "Chờ tìm kiếm (s)", 50)):
            ctk.CTkLabel(st, text=label, font=self.f_body, text_color=C_MUTED).pack(side="left", padx=(10, 4))
            e = ctk.CTkEntry(st, width=w, font=self.f_body, justify="center", fg_color=C_BG, border_color=C_BORDER)
            e.pack(side="left")
            self.e_cfg[key] = e

        # Điều khiển
        ctl = self._card(self, row=3, column=0, columnspan=2, sticky="ew", padx=18, pady=6)
        ctl.grid_columnconfigure(6, weight=1)
        self.b_login = self._btn(ctl, "🔐 Mở Zalo / Đăng nhập", self._open_zalo, color="#334155",
                                 hover="#475569", width=190)
        self.b_login.grid(row=0, column=0, padx=(14, 6), pady=(12, 6))
        self.b_start = self._btn(ctl, "🚀 Bắt đầu gửi", self._start, width=150)
        self.b_start.grid(row=0, column=1, padx=6, pady=(12, 6))
        self.b_pause = self._btn(ctl, "⏸ Tạm dừng", self._toggle_pause, color="#b45309", hover=C_WARN, width=130)
        self.b_pause.grid(row=0, column=2, padx=6, pady=(12, 6))
        self.b_stop = self._btn(ctl, "⏹ Dừng", self._stop, color="#b91c1c", hover=C_ERR, width=100)
        self.b_stop.grid(row=0, column=3, padx=6, pady=(12, 6))
        self._btn(ctl, "📋 Copy số lỗi", self._copy_failed, color="#334155", hover="#475569", width=130).grid(
            row=0, column=4, padx=6, pady=(12, 6))
        self._btn(ctl, "📁 Log", self._open_logs, color="#334155", hover="#475569", width=80).grid(
            row=0, column=5, padx=6, pady=(12, 6))

        stats = ctk.CTkFrame(ctl, fg_color="transparent")
        stats.grid(row=0, column=6, sticky="e", padx=14)
        self.lbl_stats = {}
        for key, color in (("OK", C_OK), ("UNCONFIRMED", C_WARN), ("NOT_FOUND", "#60a5fa"), ("ERROR", C_ERR)):
            lbl = ctk.CTkLabel(stats, text="", font=self.f_h, text_color=color)
            lbl.pack(side="left", padx=7)
            self.lbl_stats[key] = lbl

        self.progress = ctk.CTkProgressBar(ctl, height=10, corner_radius=5, progress_color=C_ACCENT)
        self.progress.grid(row=1, column=0, columnspan=7, sticky="ew", padx=14, pady=(4, 2))
        self.progress.set(0)
        self.lbl_status = ctk.CTkLabel(ctl, text="Sẵn sàng", font=self.f_body, text_color=C_MUTED, anchor="w")
        self.lbl_status.grid(row=2, column=0, columnspan=7, sticky="ew", padx=14, pady=(0, 10))

        # Nhật ký
        lg = self._card(self, row=4, column=0, columnspan=2, sticky="nsew", padx=18, pady=(6, 16))
        lg.grid_columnconfigure(0, weight=1)
        lg.grid_rowconfigure(0, weight=1)
        self.txt_log = ctk.CTkTextbox(lg, font=self.f_mono, fg_color=C_BG, corner_radius=10, state="disabled")
        self.txt_log.grid(row=0, column=0, sticky="nsew", padx=10, pady=10)
        for tag, color in (("info", "#c9d4e5"), ("ok", C_OK), ("warn", C_WARN), ("error", C_ERR)):
            self.txt_log.tag_config(tag, foreground=color)

        self._refresh_stats()
        self._set_busy(False)

    # ---------- Tiện ích GUI ----------
    def log(self, text: str, level: str = "info"):
        self.txt_log.configure(state="normal")
        self.txt_log.insert("end", f"[{datetime.now():%H:%M:%S}] {text}\n", level)
        self.txt_log.see("end")
        self.txt_log.configure(state="disabled")

    def _refresh_stats(self):
        for key, lbl in self.lbl_stats.items():
            lbl.configure(text=f"{STATUS_LABEL[key][0]}: {self.counts[key]}")

    def _set_busy(self, busy: bool):
        self.busy = busy
        self.b_start.configure(state="disabled" if busy else "normal")
        self.b_login.configure(state="disabled" if busy else "normal")
        self.b_pause.configure(state="normal" if busy else "disabled", text="⏸ Tạm dừng")
        self.b_stop.configure(state="normal" if busy else "disabled")
        if not busy:
            self.worker.pause_event.clear()

    def _read_cfg(self) -> SendConfig | None:
        try:
            c = SendConfig(
                delay_min=float(self.e_cfg["delay_min"].get()),
                delay_max=float(self.e_cfg["delay_max"].get()),
                batch_size=int(float(self.e_cfg["batch_size"].get())),
                batch_pause=float(self.e_cfg["batch_pause"].get()),
                search_timeout=float(self.e_cfg["search_timeout"].get()),
            )
        except ValueError:
            messagebox.showerror("Cài đặt sai", "Các ô cài đặt phải là số.")
            return None
        if c.delay_min < 0 or c.delay_max < c.delay_min or c.search_timeout <= 0:
            messagebox.showerror("Cài đặt sai", "Thời gian nghỉ 'từ' phải ≤ 'đến' và không âm.")
            return None
        return c

    def _load_config(self):
        data = {}
        if CONFIG_FILE.exists():
            try:
                data = json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                data = {}
        cfg = SendConfig(**{k: v for k, v in data.get("cfg", {}).items() if k in SendConfig.__annotations__})
        for k, e in self.e_cfg.items():
            v = getattr(cfg, k)
            e.insert(0, str(int(v)) if float(v).is_integer() else str(v))
        self.txt_msg.insert("1.0", data.get("message") or DEFAULT_MESSAGE)

    def _save_config(self, cfg: SendConfig):
        try:
            CONFIG_FILE.write_text(json.dumps({"cfg": asdict(cfg), "message": self._message()},
                                              ensure_ascii=False, indent=2), encoding="utf-8")
        except OSError:
            pass

    def _message(self) -> str:
        return self.txt_msg.get("1.0", "end").rstrip("\n")

    # ---------- Hành động ----------
    def _open_file(self):
        path = filedialog.askopenfilename(
            title="Chọn file danh sách SĐT",
            filetypes=[("Danh sách SĐT", "*.txt *.csv *.xlsx *.xlsm"), ("Tất cả", "*.*")])
        if not path:
            return
        try:
            text = read_phone_file(path)
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Lỗi đọc file", str(e))
            return
        self.txt_phones.delete("1.0", "end")
        self.txt_phones.insert("1.0", text)
        self.log(f"Đã nạp file: {os.path.basename(path)}")
        self._normalize()

    def _normalize(self) -> list[str]:
        phones, invalid, dup = parse_phones(self.txt_phones.get("1.0", "end"))
        self.txt_phones.delete("1.0", "end")
        self.txt_phones.insert("1.0", "\n".join(phones))
        self.lbl_phone_count.configure(
            text=f"Hợp lệ: {len(phones)}  •  Trùng: {dup}  •  Lỗi: {len(invalid)}",
            text_color=C_WARN if invalid else C_MUTED)
        if invalid:
            preview = ", ".join(invalid[:10]) + (" ..." if len(invalid) > 10 else "")
            self.log(f"Bỏ qua {len(invalid)} dòng không phải SĐT hợp lệ: {preview}", "warn")
        return phones

    def _preview(self):
        self.lbl_preview.configure(text=spin(self._message()) or "(trống)", text_color="#e6edf7")

    def _open_zalo(self):
        self.worker.stop_event.clear()
        self._set_busy(True)
        self.worker.submit("open")

    def _start(self):
        phones = self._normalize()
        if not phones:
            messagebox.showwarning("Thiếu dữ liệu", "Chưa có số điện thoại hợp lệ nào.")
            return
        template = self._message()
        if not template.strip():
            messagebox.showwarning("Thiếu dữ liệu", "Chưa nhập nội dung tin nhắn.")
            return
        cfg = self._read_cfg()
        if cfg is None:
            return
        if not messagebox.askyesno("Xác nhận gửi",
                                   f"Gửi tin nhắn tới {len(phones)} số điện thoại?\n\n"
                                   f"Ví dụ nội dung:\n{spin(template)}"):
            return
        self._save_config(cfg)
        self.counts = {k: 0 for k in STATUS_LABEL}
        self.failed = []
        self._refresh_stats()
        self.progress.set(0)
        self.worker.stop_event.clear()
        self.worker.pause_event.clear()
        self._set_busy(True)
        self.worker.submit("send", phones=phones, template=template, cfg=cfg)

    def _toggle_pause(self):
        if self.worker.pause_event.is_set():
            self.worker.pause_event.clear()
            self.b_pause.configure(text="⏸ Tạm dừng")
            self.log("Tiếp tục gửi.", "info")
        else:
            self.worker.pause_event.set()
            self.b_pause.configure(text="▶ Tiếp tục")
            self.log("Đã tạm dừng (sẽ dừng sau khi xong số hiện tại).", "warn")

    def _stop(self):
        self.worker.stop_event.set()
        self.worker.pause_event.clear()
        self.lbl_status.configure(text="Đang dừng...")

    def _copy_failed(self):
        if not self.failed:
            messagebox.showinfo("Copy số lỗi", "Không có số nào bị lỗi / không tìm thấy.")
            return
        self.clipboard_clear()
        self.clipboard_append("\n".join(self.failed))
        self.log(f"Đã copy {len(self.failed)} số lỗi vào clipboard (dán lại để gửi lại).", "ok")

    def _open_logs(self):
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        os.startfile(LOG_DIR)  # type: ignore[attr-defined]  # Windows

    # ---------- Nhận sự kiện từ worker ----------
    def _poll(self):
        try:
            while True:
                kind, p = self.ui_q.get_nowait()
                if kind == "log":
                    self.log(p[1], p[0])
                elif kind == "status":
                    self.lbl_status.configure(text=p)
                elif kind == "busy":
                    self._set_busy(True)
                elif kind == "idle":
                    self._set_busy(False)
                elif kind == "paused":
                    self.b_pause.configure(text="▶ Tiếp tục")
                elif kind == "login":
                    self.lbl_login.configure(text="●  Đã đăng nhập Zalo" if p else "●  Chưa kết nối Zalo",
                                             text_color=C_OK if p else C_MUTED)
                elif kind == "result":
                    self._on_result(p)
                elif kind == "done":
                    self._on_done(p)
        except queue.Empty:
            pass
        self.after(100, self._poll)

    def _on_result(self, r: dict):
        st = r["status"]
        self.counts[st] += 1
        if st in ("NOT_FOUND", "ERROR"):
            self.failed.append(r["phone"])
        self._refresh_stats()
        self.progress.set(r["i"] / r["total"])
        label, level = STATUS_LABEL[st]
        extra = f" → {r['name']}" if r["name"] else ""
        detail = f" ({r['detail']})" if r["detail"] else ""
        self.log(f"[{r['i']}/{r['total']}] {r['phone']}{extra}: {label}{detail}", level)

    def _on_done(self, p: dict):
        c = p["counts"]
        head = "⏹ ĐÃ DỪNG" if p["stopped"] else "🏁 HOÀN TẤT"
        summary = (f"{head}: gửi {c['OK']} • chưa xác nhận {c['UNCONFIRMED']} • "
                   f"không tìm thấy {c['NOT_FOUND']} • lỗi {c['ERROR']}")
        self.log(summary, "ok")
        self.log(f"File log: {p['log']}")
        self.lbl_status.configure(text=summary)

    def _on_close(self):
        if self.busy and not messagebox.askyesno("Thoát", "Đang gửi tin. Dừng và thoát?"):
            return
        self.worker.stop_event.set()
        self.worker.pause_event.clear()
        self.worker.submit("quit")
        self.worker.join(timeout=8)
        self.destroy()


def main():
    App().mainloop()


if __name__ == "__main__":
    main()
